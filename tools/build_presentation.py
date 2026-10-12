"""Build an editable deck with Norwegian notes and a matching standalone PDF."""
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import simpleSplit
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]


def build():
    data = json.loads((ROOT/'documentation/presentation.json').read_text())
    output = ROOT/'documentation/presentation'
    output.mkdir(exist_ok=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    pdf = canvas.Canvas(str(output/'DockEye.pdf'),pagesize=(960,540))
    fonts = Path('/usr/share/fonts/truetype/dejavu')
    pdfmetrics.registerFont(TTFont('DejaVu',str(fonts/'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('DejaVu-Bold',str(fonts/'DejaVuSans-Bold.ttf')))
    pdf.setTitle('DockEye — updated presentation')
    for number, item in enumerate(data['slides'],1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string('101E2B')
        title = slide.shapes.add_textbox(Inches(0.7),Inches(0.6),Inches(12),Inches(0.9)).text_frame
        title.paragraphs[0].text = item['title']
        title.paragraphs[0].font.size = Pt(32)
        title.paragraphs[0].font.bold = True
        title.paragraphs[0].font.color.rgb = RGBColor.from_string('69D6C4')
        body = slide.shapes.add_textbox(Inches(0.8),Inches(1.9),Inches(11.7),Inches(4.6)).text_frame
        body.word_wrap = True
        for index, bullet in enumerate(item['bullets']):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = '• '+bullet
            paragraph.font.size = Pt(22)
            paragraph.font.color.rgb = RGBColor.from_string('F0F4F8')
            paragraph.space_after = Pt(22)
        footer = slide.shapes.add_textbox(Inches(0.8),Inches(7),Inches(11.7),Inches(0.3)).text_frame
        footer.paragraphs[0].text = f'DockEye  |  Implementation ≠ measured accuracy                                      {number:02d}'
        footer.paragraphs[0].font.size = Pt(11)
        footer.paragraphs[0].font.color.rgb = RGBColor.from_string('9BAEBB')
        slide.notes_slide.notes_text_frame.text = item['notes']
        pdf.setFillColor(HexColor('#101E2B'))
        pdf.rect(0,0,960,540,fill=1,stroke=0)
        pdf.setFillColor(HexColor('#69D6C4'))
        pdf.setFont('DejaVu-Bold',32)
        pdf.drawString(50,468,item['title'])
        pdf.setFillColor(HexColor('#F0F4F8'))
        y = 386
        for bullet in item['bullets']:
            lines = simpleSplit('• '+bullet,'DejaVu',20,840)
            pdf.setFont('DejaVu',20)
            for line in lines:
                pdf.drawString(58,y,line)
                y -= 27
            y -= 18
        if y < 55:
            raise ValueError(f'Slide {number} overflows; shorten its text')
        pdf.setFillColor(HexColor('#9BAEBB'))
        pdf.setFont('DejaVu',11)
        pdf.drawString(58,26,'DockEye | Implementation is not measured accuracy')
        pdf.drawRightString(902,26,f'{number:02d}')
        pdf.showPage()
    presentation.save(output/'DockEye.pptx')
    pdf.save()
    print(output)


if __name__ == '__main__':
    build()
